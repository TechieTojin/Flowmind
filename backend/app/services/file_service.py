"""Upload validation and in-memory parsing of CSV / XLSX files into a DataFrame.

Nothing here writes uploads to disk or executes uploaded content:
- CSV is decoded as text and parsed by pandas.
- XLSX is opened with openpyxl in read-only, data-only mode: cached cell values
  are read, formulas are never evaluated and macros are never loaded.
"""

import csv
import io
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any

import pandas as pd
from fastapi import UploadFile
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx"}
ZIP_MAGIC = b"PK\x03\x04"
# Guard against decompression bombs: XLSX is compressed, so 20 MB can expand massively.
MAX_XLSX_CELLS = 10_000_000
_READ_CHUNK = 1024 * 1024


class FileValidationError(Exception):
    """A client-facing upload error. `message` is safe to return to clients."""

    status_code = 422

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class FileTooLargeError(FileValidationError):
    status_code = 413


class UnsupportedFileTypeError(FileValidationError):
    status_code = 415


class UnreadableFileError(FileValidationError):
    status_code = 422


@dataclass
class ParsedTable:
    dataframe: pd.DataFrame
    file_type: str
    analyzed_sheet: str | None = None
    available_sheets: list[str] = field(default_factory=list)
    duplicate_column_names: list[str] = field(default_factory=list)


async def read_upload_limited(upload: UploadFile, max_bytes: int) -> bytes:
    """Read an upload into memory, failing fast once it exceeds `max_bytes`."""
    buffer = bytearray()
    while chunk := await upload.read(_READ_CHUNK):
        buffer.extend(chunk)
        if len(buffer) > max_bytes:
            raise FileTooLargeError(
                f"File exceeds the maximum upload size of {max_bytes // (1024 * 1024)} MB."
            )
    return bytes(buffer)


def sanitize_filename(filename: str | None) -> str:
    name = PurePath((filename or "").replace("\\", "/")).name.strip()
    return name[:255] or "upload"


def detect_file_type(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    file_type = SUPPORTED_EXTENSIONS.get(suffix)
    if file_type is None:
        raise UnsupportedFileTypeError(
            "Unsupported file type. Only .csv and .xlsx files are supported."
        )
    return file_type


def parse_file(filename: str, content: bytes) -> ParsedTable:
    """Validate and parse uploaded bytes. Raises FileValidationError subclasses."""
    file_type = detect_file_type(filename)
    if not content.strip():
        raise UnreadableFileError("The uploaded file is empty.")

    table = _parse_csv(content) if file_type == "csv" else _parse_xlsx(content)

    if table.dataframe.shape[1] == 0:
        raise UnreadableFileError("The file does not contain any columns.")
    if table.dataframe.shape[0] == 0:
        raise UnreadableFileError("The file contains column headers but no data rows.")
    return table


# ---------------------------------------------------------------- headers


def _build_headers(raw_headers: list[Any]) -> tuple[list[str], list[str]]:
    """Normalise header cells to unique strings, reporting duplicated names.

    Duplicates are renamed `name.1`, `name.2`, ... (matching pandas' convention).
    """
    names = [
        str(value).strip() if value is not None and str(value).strip() else f"Unnamed: {index}"
        for index, value in enumerate(raw_headers)
    ]
    duplicates = sorted(name for name, count in Counter(names).items() if count > 1)

    seen: Counter[str] = Counter()
    used = set(names)
    unique_names: list[str] = []
    for name in names:
        if seen[name] == 0:
            unique_names.append(name)
        else:
            candidate = f"{name}.{seen[name]}"
            while candidate in used:
                seen[name] += 1
                candidate = f"{name}.{seen[name]}"
            used.add(candidate)
            unique_names.append(candidate)
        seen[name] += 1
    return unique_names, duplicates


# ---------------------------------------------------------------- CSV


def _decode_csv(content: bytes) -> str:
    sample = content[:8192]
    if sample.startswith(ZIP_MAGIC) or b"\x00" in sample:
        raise UnreadableFileError("The file does not appear to be a text CSV file.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableFileError("Could not decode the CSV file. Please save it as UTF-8.")


def _parse_csv(content: bytes) -> ParsedTable:
    text = _decode_csv(content)
    if not text.strip():
        raise UnreadableFileError("The uploaded file is empty.")

    try:
        dataframe = pd.read_csv(io.StringIO(text), skip_blank_lines=True)
    except pd.errors.EmptyDataError as exc:
        raise UnreadableFileError("The CSV file does not contain any columns.") from exc
    except (pd.errors.ParserError, ValueError, csv.Error) as exc:
        raise UnreadableFileError(
            "The CSV file is malformed and could not be parsed (check delimiters and quoting)."
        ) from exc

    # pandas silently renames duplicate headers, so detect them from the raw header row.
    try:
        raw_header = next(csv.reader(io.StringIO(text.lstrip("\r\n"))), [])
    except csv.Error:
        raw_header = []
    headers, duplicates = _build_headers(raw_header)
    if len(headers) == dataframe.shape[1]:
        dataframe.columns = headers
    else:
        dataframe.columns = [str(column) for column in dataframe.columns]
        duplicates = []

    return ParsedTable(dataframe=dataframe, file_type="csv", duplicate_column_names=duplicates)


# ---------------------------------------------------------------- XLSX


def _is_empty_cell(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_sheet_rows(worksheet: Any) -> list[list[Any]]:
    """Return non-blank rows (empty strings normalised to None) from a worksheet."""
    rows: list[list[Any]] = []
    cells = 0
    for row in worksheet.iter_rows(values_only=True):
        values = [None if _is_empty_cell(value) else value for value in row]
        cells += len(values)
        if cells > MAX_XLSX_CELLS:
            raise FileTooLargeError("The worksheet is too large to analyze.")
        if any(value is not None for value in values):
            rows.append(values)
    return rows


def _rows_to_table(rows: list[list[Any]]) -> tuple[pd.DataFrame, list[str]]:
    width = max(
        (index + 1 for row in rows for index, value in enumerate(row) if value is not None),
        default=0,
    )
    padded = [(row + [None] * width)[:width] for row in rows]
    headers, duplicates = _build_headers(padded[0])
    dataframe = pd.DataFrame(padded[1:], columns=headers).infer_objects()
    return dataframe, duplicates


def _parse_xlsx(content: bytes) -> ParsedTable:
    if not content.startswith(ZIP_MAGIC):
        raise UnreadableFileError("The file is not a valid .xlsx workbook.")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, KeyError, XML parser errors, ...
        raise UnreadableFileError("The .xlsx workbook is malformed or unreadable.") from exc

    try:
        available_sheets = list(workbook.sheetnames)
        for worksheet in workbook.worksheets:
            rows = _read_sheet_rows(worksheet)
            if rows:
                dataframe, duplicates = _rows_to_table(rows)
                return ParsedTable(
                    dataframe=dataframe,
                    file_type="xlsx",
                    analyzed_sheet=worksheet.title,
                    available_sheets=available_sheets,
                    duplicate_column_names=duplicates,
                )
    except FileValidationError:
        raise
    except Exception as exc:
        raise UnreadableFileError("The .xlsx workbook is malformed or unreadable.") from exc
    finally:
        workbook.close()

    raise UnreadableFileError("The workbook does not contain any non-empty worksheet.")
